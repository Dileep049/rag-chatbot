import os
import re
from pathlib import Path
from typing import List
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from app.config import DATA_DIR, CATEGORIES, SUPPORTED_EXTENSIONS

class MultiCategoryDocumentLoader:
    """Loads PDF, DOCX, and TXT documents across all category folders inside data/."""
    
    def __init__(self, data_dir: str = DATA_DIR):
        self.data_dir = Path(data_dir)

    @staticmethod
    def load_document(file_path: str, category: str = "general") -> List[Document]:
        """Loads a single PDF, DOCX, or TXT file with category metadata."""
        p = Path(file_path)
        loader_inst = MultiCategoryDocumentLoader()
        return loader_inst._load_single_file(p, category=category.lower())

    def load_category_documents(self, category: str) -> List[Document]:
        cat_folder = self.data_dir / category.lower()
        if not cat_folder.exists():
            return []
            
        docs = []
        for file_path in cat_folder.glob("*.*"):
            ext = file_path.suffix.lower()
            if ext in SUPPORTED_EXTENSIONS:
                loaded = self._load_single_file(file_path, category=category.lower())
                docs.extend(loaded)
        return docs

    def load_all_documents(self) -> List[Document]:
        all_docs = []
        for cat in CATEGORIES:
            cat_docs = self.load_category_documents(cat)
            all_docs.extend(cat_docs)
        return all_docs

    def load_demo_documents(self) -> List[Document]:
        demo_folder = self.data_dir / "demo"
        if not demo_folder.exists():
            return []

        docs = []
        for file_path in demo_folder.glob("*.txt"):
            filename = file_path.name
            cat_key = filename.replace("_demo.txt", "").lower()
            if cat_key == "driving_licence":
                cat_key = "driving_license"

            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()

                raw_sections = re.split(r'(?=\n(?:##\s*)?Section:)', content)
                sec_idx = 1
                for sec in raw_sections:
                    sec_clean = sec.strip()
                    if not sec_clean or (sec_clean.startswith("DEMO KNOWLEDGE BASE") and "Section:" not in sec_clean):
                        continue

                    sec_title_match = re.search(r'(?:##\s*)?Section:\s*(.+)', sec_clean)
                    section_name = sec_title_match.group(1).strip() if sec_title_match else f"Section {sec_idx}"
                    
                    page_match = re.search(r'Page:\s*(.+)', sec_clean)
                    page_str = page_match.group(1).strip() if page_match else f"Demo Page {sec_idx}"

                    docs.append(Document(
                        page_content=sec_clean,
                        metadata={
                            "document": filename,
                            "category": cat_key,
                            "page": page_str,
                            "source": "demo",
                            "source_type": "Demo Knowledge Base",
                            "section": section_name
                        }
                    ))
                    sec_idx += 1
            except Exception as e:
                print(f"Error reading demo TXT {filename}: {e}")
        return docs

    def _load_single_file(self, file_path: Path, category: str) -> List[Document]:
        ext = file_path.suffix.lower()
        filename = file_path.name
        
        if ext == ".pdf":
            return self._load_pdf(file_path, filename, category)
        elif ext in [".docx", ".doc"]:
            return self._load_docx(file_path, filename, category)
        elif ext == ".txt":
            return self._load_txt(file_path, filename, category)
        return []

    def _load_pdf(self, file_path: Path, filename: str, category: str) -> List[Document]:
        try:
            loader = PyPDFLoader(str(file_path))
            raw_docs = loader.load()
            docs = []
            for doc in raw_docs:
                page_num = doc.metadata.get("page", 0) + 1
                doc.metadata = {
                    "document": filename,
                    "category": category,
                    "page": page_num,
                    "source": "official",
                    "section": f"Page {page_num}"
                }
                docs.append(doc)
            return docs
        except Exception as e:
            print(f"Error reading PDF {filename}: {e}")
            return []

    def _load_docx(self, file_path: Path, filename: str, category: str) -> List[Document]:
        try:
            import docx
            doc = docx.Document(str(file_path))
            docs = []
            current_section = "General Section"
            current_lines = []
            page_counter = 1

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                if p.style.name.startswith('Heading') or text.startswith(('##', 'Section')):
                    if current_lines:
                        docs.append(Document(
                            page_content="\n".join(current_lines),
                            metadata={
                                "document": filename,
                                "category": category,
                                "page": page_counter,
                                "source": "official",
                                "section": current_section
                            }
                        ))
                        page_counter += 1
                        current_lines = []
                    current_section = text
                else:
                    current_lines.append(text)

            if current_lines:
                docs.append(Document(
                    page_content="\n".join(current_lines),
                    metadata={
                        "document": filename,
                        "category": category,
                        "page": page_counter,
                        "source": "official",
                        "section": current_section
                    }
                ))
            return docs
        except Exception as e:
            print(f"Error reading DOCX {filename}: {e}")
            return []

    def _load_txt(self, file_path: Path, filename: str, category: str) -> List[Document]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()

            sections = re.split(r'(?=\n##\s+)', content)
            docs = []
            for idx, sec in enumerate(sections, 1):
                sec_clean = sec.strip()
                if not sec_clean:
                    continue
                header_match = re.match(r'##\s*(.+)', sec_clean)
                section_name = header_match.group(1).strip() if header_match else "General Section"
                
                docs.append(Document(
                    page_content=sec_clean,
                    metadata={
                        "document": filename,
                        "category": category,
                        "page": idx,
                        "source": "official",
                        "section": section_name
                    }
                ))
            return docs if docs else [Document(
                page_content=content,
                metadata={
                    "document": filename,
                    "category": category,
                    "page": 1,
                    "source": "official",
                    "section": "General"
                }
            )]
        except Exception as e:
            print(f"Error reading TXT {filename}: {e}")
            return []

# Alias for backward compatibility
DocumentLoader = MultiCategoryDocumentLoader
